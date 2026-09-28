"""Build the defence deck from the DTU template with python-pptx.

Run from the repo root:  python docs/defence/build_deck.py
Output: docs/defence/defence.pptx
"""
import copy
import zipfile
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor

TEMPLATE = "docs/defence/DTU Template 16_9 - Navy Blue EN.potx"
WORK = "docs/defence/_template_as_pptx.pptx"
OUT = "docs/defence/defence.pptx"
PICS = "docs/defence/pics/"
NAVY = RGBColor(0x03, 0x0F, 0x4F)
ORANGE = RGBColor(0xE0, 0x73, 0x1F)
INK = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x66, 0x66, 0x66)


def repack_template():
    zin = zipfile.ZipFile(TEMPLATE)
    zout = zipfile.ZipFile(WORK, "w", zipfile.ZIP_DEFLATED)
    for it in zin.infolist():
        data = zin.read(it.filename)
        if it.filename == "[Content_Types].xml":
            data = data.replace(b"presentationml.template.main+xml",
                                b"presentationml.presentation.main+xml")
        zout.writestr(it, data)
    zout.close()


repack_template()
prs = Presentation(WORK)
# drop the template's sample slides
sld = prs.slides._sldIdLst
for s in list(sld):
    prs.part.drop_rel(s.rId)
    sld.remove(s)

L = {l.name.strip(): l for l in prs.slide_masters[0].slide_layouts}
W, H = prs.slide_width, prs.slide_height


def title(slide, text, size=28):
    t = slide.shapes.title
    t.text_frame.text = text
    for p in t.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)
    return t


def textbox(slide, x, y, w, h, lines, size=16, color=INK, bold_first=False,
            align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=6):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        if isinstance(line, tuple):
            txt, opts = line
        else:
            txt, opts = line, {}
        r = p.add_run()
        r.text = txt
        r.font.size = Pt(opts.get("size", size))
        r.font.bold = opts.get("bold", bold_first and i == 0)
        r.font.italic = opts.get("italic", False)
        r.font.color.rgb = opts.get("color", color)
    return tb


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def picture(slide, name, x, y, w=None, h=None):
    kw = {}
    if w is not None:
        kw["width"] = Inches(w)
    if h is not None:
        kw["height"] = Inches(h)
    return slide.shapes.add_picture(PICS + name, Inches(x), Inches(y), **kw)


# ---------------------------------------------------------------- 1 title
s = prs.slides.add_slide(L["Front A"])
s.shapes.title.text_frame.text = "Higher Fidelity for Reactive Machine-Learning Potentials by Relabelling"
for p in s.shapes.title.text_frame.paragraphs:
    for r in p.runs:
        r.font.size = Pt(36)
sub = s.placeholders[1]
sub.text_frame.text = "MSc thesis defence\nBenedikt Pabinger\nSupervisors: Arghya Bhowmik, Emma Christine Lei Hovmand, Andreas Burger"
for p in sub.text_frame.paragraphs:
    for r in p.runs:
        r.font.size = Pt(18)
notes(s, "Title. One sentence on what the talk is about: raising the level of theory "
         "of a reactive machine-learning potential by recomputing labels, and where that works.")

# ---------------------------------------------------------------- 2 MLIPs
s = prs.slides.add_slide(L["Two Content"])
title(s, "Machine-learned potentials make DFT cheap, but for reactions they had no data")
left, right = s.placeholders[1], s.placeholders[2]
left.text_frame.text = ""
p = left.text_frame.paragraphs[0]
p.text = "What they are used for"
p.runs[0].font.bold = True
p.runs[0].font.size = Pt(18)
for line in ["energy and forces in milliseconds instead of hours",
             "molecular dynamics",
             "screening thousands of molecules",
             "systems too large for DFT",
             "universal models: UMA, eSEN, trained on 10⁸ DFT calculations"]:
    q = left.text_frame.add_paragraph()
    q.text = line
    q.level = 1
    q.runs[0].font.size = Pt(16)
right.text_frame.text = ""
p = right.text_frame.paragraphs[0]
p.text = "The exception: reactions"
p.runs[0].font.bold = True
p.runs[0].font.size = Pt(18)
for line in ["the datasets held equilibrium structures",
             "a transition state is far from equilibrium",
             "“limited success as surrogate potentials for reaction barrier search” "
             "(Transition1x paper, 2022)",
             "and barriers matter: they enter the rate exponentially, 43 meV is a factor of five"]:
    q = right.text_frame.add_paragraph()
    q.text = line
    q.level = 1
    q.runs[0].font.size = Pt(16)
notes(s, "A trained model gives energy and forces in milliseconds instead of hours. That is why "
         "they are used for dynamics, screening, and systems too large for DFT. Reactions were "
         "the exception: the data sat at equilibrium and a transition state is far from it. "
         "The Transition1x paper says it in one sentence. Barriers enter the rate exponentially: "
         "43 meV, chemical accuracy, is a factor of five at room temperature.")

# ---------------------------------------------------------------- 3 Transition1x
s = prs.slides.add_slide(L["Kun titel"])
title(s, "Transition1x gave them the data, at a cheap level")
tiles = [("10 073", "reaction paths, every intermediate NEB geometry kept"),
         ("9.6 million", "DFT energies and forces on and around the paths"),
         ("ωB97X/6-31G(d)", "level chosen for cost and for compatibility with ANI-1x, not for barrier accuracy")]
x0, wtile, gap = 1.9, 3.2, 0.3
for i, (big, small) in enumerate(tiles):
    x = x0 + i * (wtile + gap)
    textbox(s, x, 2.1, wtile, 0.9, [(big, {"size": 34 if i < 2 else 26, "bold": True, "color": NAVY})],
            anchor=MSO_ANCHOR.BOTTOM)
    textbox(s, x, 3.05, wtile, 1.4, [small], size=15, color=INK)
textbox(s, 1.9, 4.6, 10.2, 1.6, [
    ("Better levels exist. ωB97M-V/def2-TZVP is among the most accurate DFT methods for barriers.", {}),
    ("A model trained on Transition1x is only as good as these labels.", {"bold": True}),
], size=18, space_after=10)
textbox(s, 1.9, 6.6, 10.2, 0.4, ["Schreiner, Bhowmik, Vegge, Busk, Winther, Sci. Data 2022"], size=11, color=GREY)
notes(s, "2022, from this group. Ten thousand reaction paths, every intermediate NEB geometry kept, "
         "9.6 million calculations, at wB97X/6-31G(d). The level was chosen for cost and for "
         "compatibility with ANI-1x, not for barrier accuracy. Better levels exist. "
         "TRANSITION: So we want those labels at a better level. The obvious way is to run the ten "
         "thousand path searches again at that level: hundreds of gradients per search, at a level "
         "several times more expensive. Nobody does that.")

# ---------------------------------------------------------------- 4 relabelling
s = prs.slides.add_slide(L["Kun titel"])
title(s, "A dataset is geometries plus labels. Relabelling changes the labels and keeps the geometries")
picture(s, "pic_dataset.png", 1.6, 1.9, w=6.3)
textbox(s, 8.2, 2.0, 4.6, 4.8, [
    ("The geometries come from the path search. The labels are the energy and force at each one.", {}),
    ("Only the labels are tied to the cheap level.", {"bold": True}),
    ("So recompute them: one single point per geometry at the better level. No new path search.", {}),
    ("This is how the field raises fidelity.", {}),
    ("OMol25, 2025: did it for all of Transition1x. Same geometries, new labels at "
     "ωB97M-V/def2-TZVPD, computed unrestricted, the way DFT can describe a half-broken bond.",
     {"color": ORANGE}),
], size=15, space_after=10)
notes(s, "You do not have to. A dataset is geometries plus labels. The geometries come from the "
         "path search, the labels are energy and force at each geometry. Only the labels are tied "
         "to the cheap level. So recompute them: one single point per geometry, no new path search. "
         "That is how the field raises fidelity. And in 2025 OMol25 did it at full scale: all of "
         "Transition1x, same geometries, new labels at wB97M-V/def2-TZVPD, and computed the other "
         "way DFT can be run for singlets, unrestricted, which can describe a half-broken bond.")

# ---------------------------------------------------------------- 5 two questions
s = prs.slides.add_slide(L["Kun titel"])
title(s, "Two things you can do to that picture, and a question for each")
colw, xl, xr, ytop = 4.9, 1.6, 7.4, 1.6
picture(s, "pic_rq1.png", xl, ytop, w=colw)
picture(s, "pic_rq2.png", xr, ytop, w=colw)
ytext = ytop + colw / 6.4 * 3.6 + 0.1
textbox(s, xl, ytext, colw, 2.6, [
    ("RQ1   How few labels are enough to lift the model?", {"bold": True, "size": 15, "color": NAVY}),
    ("What we did: trained MACE on Transition1x, relabelled under 1 % of it at "
     "ωB97M-V/def2-TZVP, trained a small correction head on that.", {}),
    ("Tested against DFT reference paths on 30 unseen reactions.", {}),
], size=13, space_after=5)
textbox(s, xr, ytext, colw, 2.6, [
    ("RQ2   Are the geometries still the right ones on the new surface?", {"bold": True, "size": 15, "color": ORANGE}),
    ("What we did: let the three OMol25 models search transition states of 45 Transition1x "
     "reactions, checked every result with unrestricted DFT and a stability analysis.", {}),
    ("Tested whether they do as well where the surface differs as where it does not.", {}),
], size=13, space_after=5)
notes(s, "Left: swap a few labels for better ones. How few are enough? RQ1. We trained our own "
         "MACE, relabelled under one percent, trained a correction head, tested on 30 unseen "
         "reactions against DFT reference paths. Right: swap all labels for ones from a different "
         "surface, geometries untouched. Are the geometries still the right ones? RQ2. We took "
         "the three OMol25 models, let them search 45 reactions, checked every result with "
         "unrestricted DFT and a stability analysis. Both are edits to the same dataset, and both "
         "are what the field does today. The rest of the talk is these two experiments.")

prs.save(OUT)
print("written", OUT, "slides:", len(prs.slides))
