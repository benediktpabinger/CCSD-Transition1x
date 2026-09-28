"""One slide: the two research questions as two edits of the same dataset.
Built in the DTU template so it can be copied into the main deck.

Run from the repo root:  python docs/defence/build_slide_rq.py
Output: docs/defence/slide_rq.pptx
"""
import zipfile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

TEMPLATE = "docs/defence/DTU Template 16_9 - Navy Blue EN.potx"
WORK = "docs/defence/_template_as_pptx.pptx"
OUT = "docs/defence/slide_rq.pptx"
PICS = "docs/defence/pics/"
NAVY = RGBColor(0x03, 0x0F, 0x4F)
ORANGE = RGBColor(0xE0, 0x73, 0x1F)
INK = RGBColor(0x22, 0x22, 0x22)
MID = RGBColor(0x7F, 0x7F, 0x86)
LINE = RGBColor(0xD0, 0xD0, 0xD4)

zin = zipfile.ZipFile(TEMPLATE)
zout = zipfile.ZipFile(WORK, "w", zipfile.ZIP_DEFLATED)
for it in zin.infolist():
    data = zin.read(it.filename)
    if it.filename == "[Content_Types].xml":
        data = data.replace(b"presentationml.template.main+xml", b"presentationml.presentation.main+xml")
    zout.writestr(it, data)
zout.close()
prs = Presentation(WORK)
sld = prs.slides._sldIdLst
for s in list(sld):
    prs.part.drop_rel(s.rId)
    sld.remove(s)
L = {l.name.strip(): l for l in prs.slide_masters[0].slide_layouts}


def textbox(slide, x, y, w, h, runs, size=14, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=4):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    first = True
    for para in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(space_after)
        for txt, opts in para:
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(opts.get("size", size))
            r.font.bold = opts.get("bold", False)
            r.font.italic = opts.get("italic", False)
            r.font.color.rgb = opts.get("color", INK)
    return tb


s = prs.slides.add_slide(L["Kun titel"])
s.shapes.title.text_frame.text = "Two ways to relabel Transition1x"
for p in s.shapes.title.text_frame.paragraphs:
    for r in p.runs:
        r.font.size = Pt(28)

cols = [
    dict(x=1.9, color=NAVY, tag="RQ1",
         q="Can a model be lifted to a higher level of theory by relabelling only a subset of its training data?",
         how="we relabel a subset ourselves and train a correction head on a MACE model",
         table="pic_table_rq1.png", pic="pic_rq1_3d.png"),
    dict(x=7.55, color=ORANGE, tag="RQ2",
         q="Can a model learn a surface from paths that were never relaxed on that surface?",
         how="we test the OMol25 models, trained on all of Transition1x relabelled unrestricted",
         table="pic_table_rq2.png", pic="pic_rq2_3d_v2.png"),
]
W = 5.35
for c in cols:
    x = c["x"]
    # question
    textbox(s, x, 1.65, W, 0.9, [[
        (c["tag"] + "   ", {"bold": True, "color": c["color"], "size": 15}),
        (c["q"], {"bold": True, "size": 14}),
    ]])
    # how
    textbox(s, x, 2.6, W, 0.55, [[
        ("To investigate ", {"bold": True, "size": 13, "color": c["color"]}),
        (c["how"], {"size": 13}),
    ]])
    # table (left) and 3D picture (right) side by side
    s.shapes.add_picture(PICS + c["table"], Inches(x), Inches(3.5), width=Inches(2.55))
    s.shapes.add_picture(PICS + c["pic"], Inches(x + 2.6), Inches(3.25), height=Inches(2.35))

# divider
ln = s.shapes.add_connector(1, Inches(7.4), Inches(1.7), Inches(7.4), Inches(5.6))
ln.line.color.rgb = LINE
ln.line.width = Pt(1)

# bottom line
textbox(s, 1.9, 5.9, 11.0, 0.7, [[
    ("Both are edits of the same dataset. ", {"size": 14, "bold": True}),
    ("The first changes the labels on a subset. The second changes the surface the labels come from, and the geometries stay where the old surface put them.", {"size": 14}),
]], anchor=MSO_ANCHOR.TOP)

s.notes_slide.notes_text_frame.text = (
    "Both questions start from Transition1x and both relabel it. Left: only the level of theory changes, "
    "the spin formalism stays restricted, and we do it on a subset because the question is how few labels are "
    "enough. Right: the level changes and the spin formalism changes, which is what OMol25 did for all of "
    "Transition1x. For the first that means a few new labels on the same surface. For the second the surface "
    "has a different shape, and its transition state can be where no geometry is. So the second question is "
    "not about how many labels; it is about whether the geometries are still the right ones.")
prs.save(OUT)
print("written", OUT)
