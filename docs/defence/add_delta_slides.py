# -*- coding: utf-8 -*-
"""Die Methodenstrecke zu Delta, auf dem Raster des uebrigen Foliensatzes.

Der Aufbau ist eine Leiste mit vier Schritten, die auf jeder Folie an genau
derselben Stelle steht. Auf der Uebersicht sind alle vier aktiv; auf den
Schrittfolien genau einer. Dadurch weiss der Saal jederzeit, wo er ist, ohne
dass der Faden reisst.

Die Leiste wird gerechnet, nicht kopiert: verschoebe sie sich zwischen zwei
Folien um zwei Pixel, faellt es beim Weiterschalten sofort auf.

    python add_delta.py <ein.pptx> <aus.pptx>
"""
import copy
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(0x03, 0x0F, 0x4F)
ORANGE = RGBColor(0xE0, 0x73, 0x1F)
PILL_OFF = RGBColor(0xEC, 0xEC, 0xEF)
GREY_T = RGBColor(0x7F, 0x7F, 0x86)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x33, 0x33, 0x38)

# Raster wie im uebrigen Satz
ML, MR = 1.00, 12.33
CW = MR - ML
TITLE_Y, TITLE_H = 0.50, 0.85
SPINE_Y, SPINE_H = 1.35, 0.52
BODY_TOP, BOTTOM = 2.10, 6.75

STEPS = [("1", "Base model"), ("2", "Relabel"),
         ("3", "Train the head"), ("4", "Test")]
PILL_W = 2.60
GAP = (CW - len(STEPS) * PILL_W) / (len(STEPS) - 1)


def textbox(sl, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    return tf


def run(p, text, size, bold=False, color=INK, italic=False):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return r


def title(sl, text):
    tf = textbox(sl, ML, TITLE_Y, CW, TITLE_H)
    run(tf.paragraphs[0], text, 28, True, NAVY)


def spine(sl, active=None):
    """active: Index des hervorgehobenen Schritts, oder None fuer alle."""
    for i, (num, label) in enumerate(STEPS):
        x = ML + i * (PILL_W + GAP)
        on = active is None or active == i
        sh = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x),
                                 Inches(SPINE_Y), Inches(PILL_W),
                                 Inches(SPINE_H))
        sh.adjustments[0] = 0.28
        sh.fill.solid()
        sh.fill.fore_color.rgb = NAVY if on else PILL_OFF
        sh.line.fill.background()
        sh.shadow.inherit = False
        tf = sh.text_frame
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run(p, num + "   ", 12, True, WHITE if on else GREY_T)
        run(p, label, 12, on, WHITE if on else GREY_T)
        if i + 1 < len(STEPS):
            g = textbox(sl, x + PILL_W, SPINE_Y, GAP, SPINE_H,
                        MSO_ANCHOR.MIDDLE)
            g.paragraphs[0].alignment = PP_ALIGN.CENTER
            run(g.paragraphs[0], "›", 14, True, GREY_T)


def add_slide(prs, layout, before_index, like=None):
    sl = prs.slides.add_slide(layout)
    if like is not None:
        for sh in like.shapes:
            if sh.is_placeholder and sh.placeholder_format.type in (13, 16):
                sl.shapes._spTree.append(copy.deepcopy(sh._element))
    lst = prs.slides._sldIdLst
    item = list(lst)[-1]
    lst.remove(item)
    lst.insert(before_index, item)
    return sl


def bullets(tf, items, size=13, gap=6):
    for i, (txt, bold) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        run(p, txt, size, bold, INK)


src, dst = sys.argv[1], sys.argv[2]
prs = Presentation(src)
S = list(prs.slides)
sep = next(i for i, s in enumerate(S)
           if s.slide_layout.name.strip() == "Front/Pause A")
layout = S[0].slide_layout                      # 'Logo og footers'
ARCH = "docs/defence/pics/arch_all.png"

# ------------------------------------------------------- Folie A, Uebersicht
sl = add_slide(prs, layout, sep, like=S[0])
title(sl, "How MACE+Δ is built and tested")
spine(sl, None)

# Die Architektur ist fast quadratisch; an der Hoehe ausgerichtet wird sie
# 4.45 Zoll breit, nicht 6.3. Das laesst rechts Platz fuer die Pruefung.
ah = BOTTOM - BODY_TOP
sl.shapes.add_picture(ARCH, Inches(ML), Inches(BODY_TOP),
                      height=Inches(ah))

RX = 6.10
tf = textbox(sl, RX, BODY_TOP, MR - RX, 0.4)
run(tf.paragraphs[0], "what it is", 16, True, NAVY)

tf = textbox(sl, RX, BODY_TOP + 0.45, MR - RX, 1.5)
bullets(tf, [
    ("one frozen MACE encoder, two readouts, one forward pass", False),
    ("the head learns the difference between the two levels, "
     "not the surface itself", False),
])

tf = textbox(sl, RX, BODY_TOP + 1.85, MR - RX, 0.4)
run(tf.paragraphs[0], "how it is tested", 16, True, NAVY)

tf = textbox(sl, RX, BODY_TOP + 2.30, MR - RX, 2.0)
bullets(tf, [
    ("30 reactions from the test split, 10 high / 10 mid / 10 low "
     "multireference character", False),
    ("on fixed geometries — energies, forces, barriers", False),
    ("with the model driving the NEB — the barrier and the "
     "transition state it finds itself", False),
])

prs.save(dst)
print("geschrieben:", dst, " neue Folie an Position", sep + 1)
