"""The RQ2 workflow diagram as native PowerPoint shapes (editable boxes,
arrows and text), on one slide of the DTU template.

Run from the repo root:  python docs/defence/build_workflow_shapes.py
Output: docs/defence/rq2_workflow_shapes.pptx
"""
import zipfile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from lxml import etree

TEMPLATE = "docs/defence/DTU Template 16_9 - Navy Blue EN.potx"
WORK = "docs/defence/_template_as_pptx.pptx"
OUT = "docs/defence/rq2_workflow_shapes.pptx"
INK = RGBColor(0x22, 0x22, 0x22)
GREY_FILL = RGBColor(0xEC, 0xEC, 0xEF)
GREY_EDGE = RGBColor(0x8C, 0x8C, 0x92)
LIGHT_FILL = RGBColor(0xF5, 0xF5, 0xF7)
LIGHT_EDGE = RGBColor(0xB8, 0xB8, 0xBD)
ORANGE = RGBColor(0xE0, 0x73, 0x1F)
ORANGE_FILL = RGBColor(0xFB, 0xE7, 0xD6)
ARROW = RGBColor(0x6F, 0x6F, 0x75)

zin = zipfile.ZipFile(TEMPLATE)
zout = zipfile.ZipFile(WORK, "w", zipfile.ZIP_DEFLATED)
for it in zin.infolist():
    data = zin.read(it.filename)
    if it.filename == "[Content_Types].xml":
        data = data.replace(b"presentationml.template.main+xml", b"presentationml.presentation.main+xml")
    zout.writestr(it, data)
zout.close()
prs = Presentation(WORK)
sld = prs.slides._sldIdLst
for s in list(sld):
    prs.part.drop_rel(s.rId)
    sld.remove(s)
L = {l.name.strip(): l for l in prs.slide_masters[0].slide_layouts}
s = prs.slides.add_slide(L["Kun titel"])
s.shapes.title.text_frame.text = "RQ2 – how the models are tested"
for p in s.shapes.title.text_frame.paragraphs:
    for r in p.runs:
        r.font.size = Pt(26)

# axis of the drawing: 11 x 6.2 units mapped onto the slide
X0, Y0, SX, SY = 0.55, 1.55, 1.10, 0.88   # inches per unit


def U(x, y):
    return Inches(X0 + x * SX), Inches(Y0 + (6.2 - y) * SY)


def box(x, y, w, h, title, lines, fill=GREY_FILL, edge=GREY_EDGE, title_color=INK, dashed=False, tsize=13, lsize=10):
    left, top = U(x, y + h)
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(w * SX), Inches(h * SY))
    shp.adjustments[0] = 0.08
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = edge; shp.line.width = Pt(1.5)
    if dashed:
        shp.line.dash_style = 4  # dash
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.06)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = title
    r.font.size = Pt(tsize); r.font.bold = True; r.font.color.rgb = title_color
    p.space_after = Pt(4)
    for ln in lines:
        q = tf.add_paragraph(); q.alignment = PP_ALIGN.CENTER
        r = q.add_run(); r.text = ln
        r.font.size = Pt(lsize); r.font.color.rgb = INK
    return shp


def arrow(x0, y0, x1, y1, color=ARROW):
    a, b = U(x0, y0); c, d = U(x1, y1)
    con = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, a, b, c, d)
    con.line.color.rgb = color; con.line.width = Pt(2)
    ln = con.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd")); tail.set("type", "triangle"); tail.set("w", "med"); tail.set("len", "med")
    return con


# row 1
box(0.3, 4.4, 3.0, 1.5, "45 reactions",
    ["Transition1x test split", "ranked by N_FOD: low, mid, high MR", "no unrestricted reference exists"])
arrow(3.3, 5.15, 3.9, 5.15)
box(3.9, 4.4, 3.2, 1.5, "3 OMol25 models search",
    ["UMA-S, UMA-M, eSEN, as released", "relax endpoints, CI-NEB", "same settings as the DFT reference"])
arrow(7.1, 5.15, 7.7, 5.15)
box(7.7, 4.4, 3.0, 1.5, "135 transition states",
    ["one per model and reaction", "the workflow reports success"])
# row 2
arrow(9.2, 4.4, 9.2, 3.75)
box(3.9, 2.25, 6.8, 1.5, "one DFT single point at each, OMol25 protocol",
    ["ωB97M-V/def2-TZVPD, unrestricted, plus a stability analysis",
     "no optimisation: the structure is judged where the model left it"],
    fill=ORANGE_FILL, edge=ORANGE)
# row 3
for k, (title, lines) in enumerate([
        ("residual force", ["largest force component on", "the unrestricted surface:", "is it a stationary point?"]),
        ("energy", ["the barrier the model found;", "the model's own energy and", "force error at the point"]),
        ("⟨S²⟩", ["0: closed-shell", "> 0: broken-symmetry", "82 / 53 structures"])]):
    x = 3.9 + k * 2.3
    arrow(x + 1.05, 2.25, x + 1.05, 1.75)
    box(x, 0.35, 2.1, 1.4, title, lines, fill=LIGHT_FILL, edge=LIGHT_EDGE, tsize=12, lsize=9.5)
# the test
box(0.3, 0.35, 3.0, 3.4, "the test",
    ["same metrics on both groups", "",
     "closed-shell: the surfaces coincide,", "the models are on home ground", "",
     "broken-symmetry: the unrestricted", "surface has its own transition state,", "no training geometry sat on it", "",
     "do the models do as well there?"],
    fill=RGBColor(0xFF, 0xFF, 0xFF), edge=ORANGE, title_color=ORANGE, dashed=True, lsize=9.5)
arrow(3.9, 1.05, 3.35, 1.05, color=ORANGE)

prs.save(OUT)
print("written", OUT)
