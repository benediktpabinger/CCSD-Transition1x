"""Conclusion and outlook, two slides, in the DTU template (editable text).

Run from the repo root:  python docs/defence/build_slides_conclusion.py
Output: docs/defence/slides_conclusion.pptx
"""
import zipfile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_ANCHOR
from pptx.dml.color import RGBColor

TEMPLATE = "docs/defence/DTU Template 16_9 - Navy Blue EN.potx"
WORK = "docs/defence/_template_as_pptx.pptx"
OUT = "docs/defence/slides_conclusion.pptx"
NAVY = RGBColor(0x03, 0x0F, 0x4F)
INK = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x7F, 0x7F, 0x86)
LINE = RGBColor(0xD0, 0xD0, 0xD4)

zin = zipfile.ZipFile(TEMPLATE)
zout = zipfile.ZipFile(WORK, "w", zipfile.ZIP_DEFLATED)
for it in zin.infolist():
    data = zin.read(it.filename)
    if it.filename == "[Content_Types].xml":
        data = data.replace(b"presentationml.template.main+xml", b"presentationml.presentation.main+xml")
    zout.writestr(it, data)
zout.close()
prs = Presentation(WORK)
sld = prs.slides._sldIdLst
for s in list(sld):
    prs.part.drop_rel(s.rId)
    sld.remove(s)
L = {l.name.strip(): l for l in prs.slide_masters[0].slide_layouts}


def new_slide(title):
    s = prs.slides.add_slide(L["Kun titel"])
    s.shapes.title.text_frame.text = title
    for p in s.shapes.title.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(28); r.font.color.rgb = NAVY
    return s


def textbox(slide, x, y, w, h, paras, size=14, space_after=8):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space_after)
        for txt, o in runs:
            r = p.add_run(); r.text = txt
            r.font.size = Pt(o.get("size", size)); r.font.bold = o.get("bold", False)
            r.font.italic = o.get("italic", False); r.font.color.rgb = o.get("color", INK)
    return tb


def b(txt): return (txt, {"bold": True})
def n(txt): return (txt, {})
DOT = "•  "

# ------------------------------------------------------------ 1 conclusion
s = new_slide("Conclusion")
textbox(s, 0.6, 1.6, 5.9, 5.3, [
    [("RQ1", {"bold": True, "size": 16, "color": NAVY})],
    [("Can a model be lifted to a higher level of theory by relabelling only a subset of its training data?", {"italic": True, "size": 12, "color": GREY})],
    [b("Yes"), n(", as long as the spin formalism is not changed with the labels")],
    [n(DOT + "Under 1 % relabelled: force error to a third in every tier, energies improve in every tier, geometries as good as MACE's")],
    [n(DOT + "Barriers: error drops to a third at low MR; at mid and high MR the correction overshoots, mostly the base model's error")],
    [n(DOT + "The correction is bounded by the base model: 42 meV of MACE's own error pass through unchanged")],
], size=13)
textbox(s, 6.9, 1.6, 5.9, 5.3, [
    [("RQ2", {"bold": True, "size": 16, "color": NAVY})],
    [("Can a model learn a surface from paths that were never relaxed on that surface?", {"italic": True, "size": 12, "color": GREY})],
    [b("Yes when only the level of theory changes. No when the spin formalism changes as well")],
    [n(DOT + "At broken-symmetry transition states the models deliver structures that are not stationary points: DFT residual forces 2.5 times larger, while their own forces report convergence")],
    [n(DOT + "Force error two to three times larger, energies within chemical accuracy: the barrier error is a geometry error")],
    [n(DOT + "Part of this is multireference character itself; at equal MR character a factor of about 1.6 remains")],
    [n(DOT + "Two candidate causes: the training data, never relaxed on the unrestricted surface, and the shape of that surface. The data do not decide")],
], size=13)
ln = s.shapes.add_connector(1, Inches(6.7), Inches(1.7), Inches(6.7), Inches(6.7))
ln.line.color.rgb = LINE; ln.line.width = Pt(1)

# ------------------------------------------------------------ 2 implications
s = new_slide("Implications and outlook")
textbox(s, 1.2, 1.7, 11.2, 5.2, [
    [b("A test that needs no reference.  "), n("One unrestricted DFT single point with a stability analysis at the structure a model delivers. It shows for any model and any reaction whether a stationary point of the unrestricted surface was found")],
    [b("In practice.  "), n("Run it at every transition state that is trusted further: as a training point, a reference or a starting point")],
    [b("Benchmarks.  "), n("None tests models at broken-symmetry transition states; reported success rates were measured on restricted references")],
    [b("Future work.  "), n("Fine-tune on paths relaxed on the unrestricted surface. If the training data are the cause the errors disappear; if the shape of the surface is the cause they remain. The same paths give a benchmark of broken-symmetry saddles")],
], size=16, space_after=18)
prs.save(OUT)
print("written", OUT)
